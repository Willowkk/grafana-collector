from datetime import datetime
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from grafana_collector import exporter
from grafana_collector.exporter import export_run
from grafana_collector.storage import Store


def make_plan():
    panels = []
    for i, title in [(1, "Through/put"), (2, "Through/put"), (3, "Errors"), (4, "Empty")]:
        query = {"key": f"q{i}", "panel_id": i, "ref_id": "A", "kind": "opentsdb",
                 "route": "/api/query", "query": {"metric": "example"}, "interval_ms": 1000,
                 "alias": "same", "unit": "bytes", "hidden": False, "dependencies": [], "metadata": {}, "error": None}
        panels.append({"id": i, "title": title, "group": "Overview", "queries": [query], "transformations": [], "metadata": {}})
    return {"schema_version": 1, "source_url": "https://example.test/d/test", "dashboard": {"uid": "test", "version": 9},
            "variables": {"filesystem": "test"}, "panels": panels, "from_ms": 1000,
            "to_ms": 4000, "mode": "fetch", "started_ms": 11_000, "metadata": {}}


def series(host, points, name="same"):
    return {"name": name, "metric": "example", "ref_id": "A", "labels": {"host": host}, "unit": "bytes", "points": points}


@pytest.fixture
def store(tmp_path):
    with Store(tmp_path / "run") as store:
        store.initialize(make_plan())
        yield store


def load_manifest(path):
    return json.loads(path.read_text())


def test_dates_numeric_outer_join_collision_names_and_failure_status(store, tmp_path):
    samples = [series("a", [[1000, 0], [3000, 3.125]]), series("b", [[2000, 2], [3000, None]])]
    store.record_result(make_plan()["panels"][0]["queries"][0], 1000, 4000, samples, {"data": "source"})
    store.record_display(1, samples, 1000, 4000)
    store.record_display(2, [series("x", [[1000, 9]], "=HYPERLINK(\"evil\")")], 1000, 4000)
    store.record_panel_status(3, "failed", "HTTP 503")
    store.record_display(4, [], 1000, 4000)
    out = tmp_path / "export"
    manifest = load_manifest(export_run(store, out))
    panels = {p["panel_id"]: p for p in manifest["panels"]}
    assert panels[3]["status"] == "failed" and panels[3]["files"] == []
    assert panels[4]["status"] == "empty" and len(panels[4]["files"]) == 1
    assert panels[1]["files"][0]["path"] != panels[2]["files"][0]["path"]
    assert all(c["unit"] == "bytes" for c in panels[1]["columns"])
    assert len({c["column_name"] for c in panels[1]["columns"]}) == 2
    book = load_workbook(out / panels[1]["files"][0]["path"])
    assert book.sheetnames == [panels[1]["files"][0]["sheet"]]
    ws = book.active
    assert ws["A1"].value == "Time"
    assert ws["A2"].value == datetime(1970, 1, 1, 8, 0, 1)
    assert ws["A2"].is_date and ws["A2"].data_type == "d"
    assert ws["A2"].number_format == "yyyy-mm-dd h:mm:ss"
    assert ws["A1"].font.bold and ws["A1"].fill.patternType is None
    assert not ws["A1"].alignment.wrap_text
    assert ws.freeze_panes is None and ws.auto_filter.ref is None
    assert ws.sheet_format.defaultRowHeight == 15
    assert ws.column_dimensions["A"].width == 20
    assert ws.column_dimensions["B"].width >= 18
    cols_by_host = {c["labels"]["host"]: i + 2 for i, c in enumerate(panels[1]["columns"])}
    assert "B" in ws.cell(2, cols_by_host["a"]).number_format
    assert [ws.cell(row, cols_by_host["a"]).value for row in range(2, 5)] == [0, None, 3.125]
    assert [ws.cell(row, cols_by_host["b"]).value for row in range(2, 5)] == [None, 2, None]
    book.close()
    book = load_workbook(out / panels[2]["files"][0]["path"])
    assert book.active["B1"].data_type == "s"
    assert book.active["B1"].value.startswith("=HYPERLINK")
    book.close()
    assert manifest["query_statuses"][0]["attempts"][0]["raw_path"].endswith(".json.gz")


def test_split_rows_and_columns_keeps_exact_mapping(store, tmp_path):
    samples = [series(str(i), [[1000, i], [2000, i + 1], [3000, i + 2], [4000, i + 3]], f"curve {i}") for i in range(5)]
    store.record_display(1, samples, 1000, 4000)
    out = tmp_path / "split"
    manifest = load_manifest(export_run(store, out, panel_ids=[1], max_rows=3, max_columns=3))
    files = manifest["panels"][0]["files"]
    assert len(files) == 6
    seen = {}
    for file in files:
        book = load_workbook(out / file["path"])
        ws = book.active
        assert ws.max_row <= 3 and ws.max_column <= 3
        for row in range(2, ws.max_row + 1):
            timestamp = int((ws.cell(row, 1).value - datetime(1970, 1, 1, 8)).total_seconds() * 1000)
            for col, sid in enumerate(file["series_ids"], 2):
                assert (sid, timestamp) not in seen
                seen[sid, timestamp] = ws.cell(row, col).value
        book.close()
    assert len(seen) == 20


def test_offline_reexport_filters_and_original_generation_unchanged(store, tmp_path):
    store.record_display(1, [series("a", [[1000, 1], [2000, 2], [3000, 3]])], 1000, 3000)
    out = tmp_path / "export"
    original = load_manifest(export_run(store, out, panel_ids=[1]))
    original_path = out / original["panels"][0]["files"][0]["path"]
    original_bytes = original_path.read_bytes()
    filtered = load_manifest(export_run(store, out, panel_ids=[1], from_ms=2000, to_ms=2000))
    assert original["generation"] != filtered["generation"]
    assert original_path.read_bytes() == original_bytes
    assert filtered["panels"][0]["selected_point_count"] == 1
    assert filtered["panels"][0]["files"][0]["row_count"] == 1


def test_failed_export_leaves_previous_manifest_and_files_intact(store, tmp_path, monkeypatch):
    store.record_display(1, [series("a", [[1000, 1]])], 1000, 3000)
    out = tmp_path / "export"
    path = export_run(store, out, panel_ids=[1])
    previous = path.read_bytes()
    previous_generations = set((out / "exports").iterdir())
    def broken(*args):
        raise OSError("disk full")
    monkeypatch.setattr(exporter, "_write_workbook", broken)
    with pytest.raises(OSError, match="disk full"):
        export_run(store, out, panel_ids=[1])
    assert path.read_bytes() == previous
    assert set((out / "exports").iterdir()) == previous_generations


def test_partial_history_can_export_without_hiding_failure(store, tmp_path):
    store.record_display(1, [series("a", [[1000, 1]])], 1000, 3000)
    store.record_panel_status(1, "failed", "Authentication expired")
    manifest = load_manifest(export_run(store, tmp_path / "out", panel_ids=[1]))
    entry = manifest["panels"][0]
    assert entry["status"] == "partial"
    assert entry["errors"] == ["Authentication expired"]
    assert len(entry["files"]) == 1


def test_unknown_panel_and_invalid_ranges_are_errors(store, tmp_path):
    with pytest.raises(ValueError, match="Unknown panel"):
        export_run(store, tmp_path / "a", panel_ids=[999])
    with pytest.raises(ValueError, match="precedes"):
        export_run(store, tmp_path / "a", from_ms=3000, to_ms=1000)
    with pytest.raises(ValueError, match="limits"):
        export_run(store, tmp_path / "a", max_columns=1)


def test_export_uses_one_database_snapshot_while_writer_advances(store, tmp_path, monkeypatch):
    store.record_display(1, [series("a", [[1000, 1]])], 1000, 3000)
    store.record_display(2, [series("a", [[1000, 2]])], 1000, 3000)
    original_writer = exporter._write_workbook
    advanced = False
    def update_during_export(*args):
        nonlocal advanced
        if not advanced:
            with Store(store.run_dir) as collector:
                collector.record_display(2, [series("a", [[1000, 99], [2000, 100]])], 1000, 3000)
            advanced = True
        return original_writer(*args)
    monkeypatch.setattr(exporter, "_write_workbook", update_during_export)
    out = tmp_path / "snapshot"
    manifest = load_manifest(export_run(store, out, panel_ids=[1, 2]))
    second = manifest["panels"][1]
    assert second["selected_point_count"] == 1
    book = load_workbook(out / second["files"][0]["path"])
    assert book.active["B2"].value == 2
    book.close()
    assert store.get_display_series(2)[0]["points"] == [[1000, 99], [2000, 100]]


def test_uncollected_filter_is_not_a_healthy_empty_file(store, tmp_path):
    store.record_display(1, [series("a", [[1000, 1]])], 1000, 4000)
    manifest = load_manifest(export_run(store, tmp_path / "out", panel_ids=[1], from_ms=5000, to_ms=7000))
    entry = manifest["panels"][0]
    assert entry["status"] == "uncollected"
    assert entry["collection_status"] == "success"
    assert entry["range_status"] == "uncollected"
    assert entry["files"] == []


def test_partial_coverage_is_explicit_even_when_confirmed_part_is_empty(store, tmp_path):
    store.record_display(1, [], 1000, 2000)
    store.record_display(1, [], 3000, 4000)
    manifest = load_manifest(export_run(store, tmp_path / "out", panel_ids=[1], from_ms=1000, to_ms=4000))
    entry = manifest["panels"][0]
    assert entry["status"] == "partial"
    assert entry["display_coverage"]["intervals"] == [[1000, 2000], [3000, 4000]]
    assert entry["files"] == []
    manifest = load_manifest(export_run(store, tmp_path / "out", panel_ids=[1], from_ms=1100, to_ms=1900))
    entry = manifest["panels"][0]
    assert entry["status"] == "empty"
    assert entry["range_status"] == "covered"
    assert len(entry["files"]) == 1


@pytest.mark.parametrize("bounds", [{"to_ms": 500}, {"from_ms": 5000}])
def test_one_sided_outside_filter_does_not_snap_to_collected_boundary(store, tmp_path, bounds):
    store.record_display(1, [], 1000, 4000)
    manifest = load_manifest(export_run(store, tmp_path / "out", panel_ids=[1], **bounds))
    entry = manifest["panels"][0]
    assert entry["range_status"] == "uncollected"
    assert entry["status"] == "uncollected"
    assert entry["files"] == []


def test_reference_filenames_do_not_overwrite_cleaned_truncated_or_case_collisions(tmp_path):
    plan = make_plan()
    titles = ["Throughput", "Same", "same", "bad/name", "bad:name", "x" * 110 + "one",
              "x" * 110 + "two", "Same [2]", "[Invalid]:sheet/name?*\\", "中文" * 70]
    plan["panels"] = [{"id": i, "title": title, "group": "Group", "queries": [],
                       "transformations": [], "metadata": {}} for i, title in enumerate(titles, 1)]
    with Store(tmp_path / "names") as store:
        store.initialize(plan)
        for panel in plan["panels"]:
            store.record_display(panel["id"], [series("a", [[1000, panel["id"]]])], 1000, 4000)
        out = tmp_path / "full"
        manifest = load_manifest(export_run(store, out))
        files = {p["panel_id"]: p["files"][0] for p in manifest["panels"]}
        assert Path(files[1]["path"]).name == "Throughput-data-export.xlsx"
        assert Path(files[2]["path"]).name == "Same [2]-data-export.xlsx"
        assert Path(files[3]["path"]).name == "same [3]-data-export.xlsx"
        assert Path(files[4]["path"]).name == "bad_name [4]-data-export.xlsx"
        assert Path(files[5]["path"]).name == "bad_name [5]-data-export.xlsx"
        assert "[6]" in Path(files[6]["path"]).name and "[7]" in Path(files[7]["path"]).name
        assert Path(files[8]["path"]).name == "Same [2] [8]-data-export.xlsx"
        names = [Path(f["path"]).name for f in files.values()]
        assert len({name.casefold() for name in names}) == len(titles)
        assert all(len(name.encode("utf-8")) <= 255 for name in names)
        for file in files.values():
            book = load_workbook(out / file["path"])
            assert book.sheetnames == [file["sheet"]]
            assert 1 <= len(file["sheet"]) <= 31
            assert not any(c in file["sheet"] for c in '[]:*?/\\')
            book.close()
        subset = load_manifest(export_run(store, tmp_path / "subset", panel_ids=[2, 6]))
        for panel in subset["panels"]:
            assert Path(panel["files"][0]["path"]).name == Path(files[panel["panel_id"]]["path"]).name


def test_millisecond_timestamps_remain_distinct_and_visible_in_reference_date_format(store, tmp_path):
    store.record_display(1, [series("a", [[1000, 1], [1123, 2], [1999, 3], [2000, 4]])], 1000, 4000)
    out = tmp_path / "milliseconds"
    manifest = load_manifest(export_run(store, out, panel_ids=[1]))
    book = load_workbook(out / manifest["panels"][0]["files"][0]["path"])
    sheet = book.active
    assert [sheet.cell(row, 1).value for row in range(2, 6)] == [
        datetime(1970, 1, 1, 8, 0, 1), datetime(1970, 1, 1, 8, 0, 1, 123000),
        datetime(1970, 1, 1, 8, 0, 1, 999000), datetime(1970, 1, 1, 8, 0, 2)]
    assert [sheet.cell(row, 1).number_format for row in range(2, 6)] == [
        "yyyy-mm-dd h:mm:ss", "yyyy-mm-dd h:mm:ss.000", "yyyy-mm-dd h:mm:ss.000", "yyyy-mm-dd h:mm:ss"]
    assert [sheet.cell(row, 2).value for row in range(2, 6)] == [1, 2, 3, 4]
    book.close()


def test_panel_units_are_shared_across_columns_and_splits_without_changing_raw_data(store, tmp_path):
    large = series("a", [[1000, 1.5 * 1024**3], [2000, None]], "large")
    small = series("b", [[1000, 300 * 1024**2], [2000, 0]], "small")
    large["unit"] = small["unit"] = "binBps"
    store.record_display(1, [large, small], 1000, 4000)
    before = store.get_display_series(1)
    out = tmp_path / "units"
    manifest = load_manifest(export_run(store, out, panel_ids=[1], max_columns=2, max_rows=2))
    panel = manifest["panels"][0]
    assert len(panel["files"]) == 4
    columns = {c["series_id"]: c for c in panel["columns"]}
    for column in columns.values():
        assert column["unit"] == "binBps"  # Original semantic unit is preserved.
        assert column["source_unit"] == "binBps"
        assert column["display_unit"] == "GiB/s"
        assert column["scale"] == 1 / 1024**3
    restored = []
    for file in panel["files"]:
        book = load_workbook(out / file["path"])
        ws = book.active
        column = columns[file["series_ids"][0]]
        if ws["B2"].value is not None:
            assert "GiB/s" in ws["B2"].number_format
            assert ws["B2"].data_type == "n"
            restored.append(ws["B2"].value / column["scale"])
        assert ws.column_dimensions["B"].width >= 18
        book.close()
    assert sorted(restored) == [0, 300 * 1024**2, 1.5 * 1024**3]
    assert store.get_display_series(1) == before


def test_compatible_source_units_share_scale_family_and_percent_stays_numeric(store, tmp_path):
    byte = series("bytes", [[1000, 1024**2]], "byte count")
    kilo = series("kbytes", [[1000, 1024]], "kilobyte count")
    rate = series("rate", [[1000, 2 * 1024**3]], "throughput")
    ratio = series("ratio", [[1000, 0.0125]], "success rate")
    byte["unit"], kilo["unit"], rate["unit"], ratio["unit"] = "bytes", "kbytes", "binBps", "percentunit"
    store.record_display(1, [byte, kilo, rate, ratio], 1000, 4000)
    out = tmp_path / "mixed-units"
    manifest = load_manifest(export_run(store, out, panel_ids=[1]))
    panel = manifest["panels"][0]
    columns = {c["labels"]["host"]: (index + 2, c) for index, c in enumerate(panel["columns"])}
    book = load_workbook(out / panel["files"][0]["path"])
    ws = book.active
    for host in ("bytes", "kbytes"):
        index, column = columns[host]
        assert column["display_unit"] == "MiB"
        assert ws.cell(2, index).value == 1
        assert "MiB" in ws.cell(2, index).number_format
    assert columns["bytes"][1]["scale"] == 1 / 1024**2
    assert columns["kbytes"][1]["scale"] == 1 / 1024
    index, rate_column = columns["rate"]
    assert rate_column["display_unit"] == "GiB/s"
    assert ws.cell(2, index).value == 2
    index, percent_column = columns["ratio"]
    assert percent_column["source_unit"] == "percentunit"
    assert percent_column["scale"] == 1
    assert ws.cell(2, index).value == 0.0125
    assert ws.cell(2, index).data_type == "n" and "%" in ws.cell(2, index).number_format
    book.close()


def test_mixed_time_sources_and_tiny_throughput_remain_numeric_and_visible(store, tmp_path):
    micro = series("micro", [[1000, 1000000]], "microseconds")
    milli = series("milli", [[1000, 1000]], "milliseconds")
    rate = series("rate", [[1000, 2 * 1024**3], [2000, 0.1666]], "rate")
    micro["unit"], milli["unit"], rate["unit"] = "µs", "ms", "binBps"
    store.record_display(1, [micro, milli, rate], 1000, 4000)
    out = tmp_path / "time-and-small-values"
    manifest = load_manifest(export_run(store, out, panel_ids=[1]))
    panel = manifest["panels"][0]
    columns = {c["labels"]["host"]: (index + 2, c) for index, c in enumerate(panel["columns"])}
    book = load_workbook(out / panel["files"][0]["path"])
    ws = book.active
    assert columns["micro"][1]["display_unit"] == columns["milli"][1]["display_unit"]
    micro_value = ws.cell(2, columns["micro"][0]).value
    milli_value = ws.cell(2, columns["milli"][0]).value
    assert micro_value == milli_value
    assert micro_value / columns["micro"][1]["scale"] == pytest.approx(1000000)
    assert milli_value / columns["milli"][1]["scale"] == pytest.approx(1000)
    index, column = columns["rate"]
    tiny = ws.cell(3, index)
    assert tiny.data_type == "n" and tiny.value > 0
    assert tiny.value / column["scale"] == pytest.approx(0.1666)
    assert "E" in tiny.number_format and "GiB/s" in tiny.number_format
    book.close()
