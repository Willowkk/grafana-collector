"""Offline, atomic Excel snapshots of collected display data."""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import json
import hashlib
import math
import os
from pathlib import Path
import re
import shutil
import uuid
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .units import build_unit_plan, convert_value, number_format_for_value


EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
_SHANGHAI = ZoneInfo("Asia/Shanghai")


def _filename(value, fallback="unnamed"):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")
    value = value[:100]
    # Leave room for stable IDs, split suffixes and the extension on filesystems
    # whose component limit is in UTF-8 bytes rather than characters.
    while len(value.encode("utf-8")) > 180:
        value = value[:-1]
    return value.rstrip(" .") or fallback


def _panel_filenames(panels):
    """Names stay stable when exporting a subset of the frozen dashboard."""
    bases = {p["id"]: _filename(p["title"]) for p in panels}
    counts = Counter((p.get("group") or "Ungrouped", bases[p["id"]].casefold()) for p in panels)
    needs_id = {p["id"] for p in panels if bases[p["id"]] != str(p["title"])
                or counts[(p.get("group") or "Ungrouped", bases[p["id"]].casefold())] > 1}
    while True:
        names = {p["id"]: bases[p["id"]] + (f" [{p['id']}]" if p["id"] in needs_id else "")
                 for p in panels}
        buckets = {}
        for p in panels:
            buckets.setdefault((p.get("group") or "Ungrouped", names[p["id"]].casefold()), []).append(p["id"])
        collisions = {pid for ids in buckets.values() if len(ids) > 1 for pid in ids}
        if not collisions:
            return {pid: name + "-data-export" for pid, name in names.items()}
        # Handles a literal title such as "Throughput [166]" colliding with
        # another panel's generated ID suffix without overwriting either file.
        new_ids = collisions - needs_id
        if not new_ids:
            raise ValueError("Panel IDs and titles do not form unique export paths")
        needs_id.update(new_ids)


def _sheet_name(title):
    title = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", str(title)).strip(" '")
    return ((title or "Panel") + "-data")[:31].rstrip(" '")


def _column_names(series):
    names = [str(s.get("name") or s.get("metric") or s.get("ref_id") or "series") for s in series]
    counts = Counter(names)
    used = {"Time"}
    result = []
    for name, item in zip(names, series):
        suffix = item["series_id"][:10]
        candidate = f"{name} [{suffix}]" if counts[name] > 1 or name in used else name
        if candidate in used:
            candidate = f"{name} [{item['series_id']}]"
        # Excel limits cell text, independently of sheet/file names.
        if len(candidate) > 32_767:
            candidate = candidate[:32_680] + f" [{item['series_id']}]"
        used.add(candidate)
        result.append(candidate)
    return result


def _write_workbook(path, series, legends, timestamps, sheet_name, unit_specs):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = True
    sheet.sheet_format.defaultRowHeight = 15
    sheet.column_dimensions["A"].width = 20
    labels = ["Time", *legends]
    cells = []
    for index, label in enumerate(labels, 1):
        cell = WriteOnlyCell(sheet, value=label)
        # Legends come from remote data and must remain strings, not formulas.
        cell.data_type = "s"
        cell.font = Font(bold=True)
        cells.append(cell)
        if index > 1:
            sheet.column_dimensions[get_column_letter(index)].width = max(18, unit_specs[index - 2]["min_column_width"])
    sheet.append(cells)
    mappings = [dict(item.get("points", [])) for item in series]
    for timestamp in timestamps:
        value = datetime.fromtimestamp(timestamp / 1000, tz=timezone.utc).astimezone(_SHANGHAI).replace(tzinfo=None)
        time_cell = WriteOnlyCell(sheet, value=value)
        time_cell.number_format = "yyyy-mm-dd h:mm:ss.000" if timestamp % 1000 else "yyyy-mm-dd h:mm:ss"
        row = [time_cell]
        for mapping, unit_spec in zip(mappings, unit_specs):
            point = mapping.get(timestamp)
            if isinstance(point, float) and not math.isfinite(point):
                point = None
            if point is None:
                row.append(None)
            else:
                # Store remains in the source unit. The workbook uses a fixed
                # per-panel unit scale and still contains ordinary numeric cells.
                scaled = convert_value(point, unit_spec)
                cell = WriteOnlyCell(sheet, value=scaled)
                cell.number_format = number_format_for_value(scaled, unit_spec)
                row.append(cell)
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def export_run(store, out_dir, *, from_ms=None, to_ms=None, panel_ids=None,
               max_rows=EXCEL_MAX_ROWS, max_columns=EXCEL_MAX_COLUMNS):
    with store.read_snapshot():
        return _export_run(store, out_dir, from_ms=from_ms, to_ms=to_ms,
                           panel_ids=panel_ids, max_rows=max_rows, max_columns=max_columns)


def _export_run(store, out_dir, *, from_ms, to_ms, panel_ids, max_rows, max_columns):
    """Write immutable files, then atomically publish a manifest referencing them.

    Previous generations stay valid if export fails or another reader still has
    the old manifest. ``max_rows`` includes the header and ``max_columns`` the
    Time column; smaller limits are useful for bounded integration tests.
    """
    if not 2 <= max_rows <= EXCEL_MAX_ROWS or not 2 <= max_columns <= EXCEL_MAX_COLUMNS:
        raise ValueError("Excel limits require 2+ rows/columns and cannot exceed Excel's limits")
    if from_ms is not None and to_ms is not None and to_ms < from_ms:
        raise ValueError("Export end precedes start")
    out_dir = Path(out_dir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    plan = store.load_plan()
    selected_ids = set(panel_ids) if panel_ids is not None else {p["id"] for p in plan["panels"]}
    missing = selected_ids - {p["id"] for p in plan["panels"]}
    if missing:
        raise ValueError(f"Unknown panel IDs: {sorted(missing)}")
    panels = [p for p in plan["panels"] if p["id"] in selected_ids]
    panel_filenames = _panel_filenames(plan["panels"])
    generation = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S") + "-" + uuid.uuid4().hex[:12]
    exports = out_dir / "exports"
    exports.mkdir(exist_ok=True)
    staging = exports / ("." + generation + ".tmp")
    staging.mkdir()
    destination = exports / generation
    final_manifest = out_dir / "manifest.json"
    temporary_manifest = out_dir / (".manifest-" + uuid.uuid4().hex + ".tmp")
    statuses = {p["panel_id"]: p for p in store.panel_statuses()}
    query_statuses = [q for q in store.query_statuses() if q["panel_id"] in selected_ids]
    groups = list(dict.fromkeys(p.get("group") or "Ungrouped" for p in panels))
    group_names = {}
    seen_groups = set()
    for group in groups:
        name = _filename(group)
        if name.casefold() in seen_groups:
            suffix = hashlib.sha256(group.encode()).hexdigest()[:12]
            name += f" [{suffix}]"
            while name.casefold() in seen_groups:
                name += "_"
        seen_groups.add(name.casefold())
        group_names[group] = name
    manifest = {
        "schema_version": 1,
        "generation": generation,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_url": plan["source_url"],
        "source_run": str(store.run_dir),
        "dashboard": plan["dashboard"],
        "variables": plan["variables"],
        "mode": plan["mode"],
        "timezone": "Asia/Shanghai",
        "requested_range": {"from_ms": from_ms, "to_ms": to_ms},
        "dataset_start_ms": plan["from_ms"],
        "sampling": plan.get("metadata", {}),
        "precision": {
            "storage": "Source JSON numeric values; no decimal rounding in SQLite",
            "excel": "Numeric cells follow Excel floating-point precision (approximately 15 significant digits)",
            "conversion": "Excel numeric value = source value * column.scale; no pre-rounding; percentunit retains its 0..1 value with Excel percent formatting",
            "display": "Compatible units share a fixed scale per panel and selected export window; number formatting adjusts precision without changing numeric values",
            "missing": "Blank cells; no interpolation or zero filling",
        },
        "raw_provenance": None,
        "query_statuses": query_statuses,
        "panels": [],
    }
    try:
        for panel in panels:
            status = statuses[panel["id"]]
            attempted_ends = [attempt["end_ms"] for q in query_statuses if q["panel_id"] == panel["id"] for attempt in q["attempts"]]
            displayed_end = store.connection.execute("SELECT MAX(end_ms) FROM display_updates WHERE panel_id=?", (panel["id"],)).fetchone()[0]
            coverage_start = int(from_ms) if from_ms is not None else int(plan["from_ms"])
            coverage_end = int(to_ms) if to_ms is not None else max([int(plan["to_ms"]), displayed_end or int(plan["to_ms"]), *attempted_ends])
            if coverage_end < coverage_start:
                # A one-sided filter can be wholly outside the dataset. Do not
                # clamp it onto a collected boundary and invent a healthy empty.
                coverage = {"status": "uncollected", "from_ms": from_ms, "to_ms": to_ms,
                            "intervals": [], "reason": "Filter does not intersect this panel's dataset range"}
            else:
                coverage = store.display_coverage(panel["id"], coverage_start, coverage_end)
            series = store.get_display_series(panel["id"], from_ms, to_ms)
            series.sort(key=lambda s: (str(s.get("name", "")), s["series_id"]))
            legends = _column_names(series)
            timestamps = sorted({int(p[0]) for s in series for p in s.get("points", [])})
            unit_plan = build_unit_plan(series)
            entry = {
                **status,
                "collection_status": status["status"],
                "range_status": coverage["status"],
                "display_coverage": coverage,
                "query_definitions": panel["queries"],
                "transformations": panel.get("transformations", []),
                "metadata": panel.get("metadata", {}),
                "selected_point_count": sum(len(s.get("points", [])) for s in series),
                "selected_time_range": {"from_ms": timestamps[0] if timestamps else None, "to_ms": timestamps[-1] if timestamps else None},
                "columns": [{**{k: v for k, v in item.items() if k != "points"}, "column_name": legend,
                             **unit_plan[item["series_id"]]}
                            for item, legend in zip(series, legends)],
                "files": [],
            }
            if status["status"] in {"success", "empty"} and coverage["status"] != "covered":
                entry["status"] = coverage["status"]
            # Previously valid data can be exported after a failed later attempt,
            # but the failure remains explicit in the manifest.
            valid_empty = not timestamps and status["status"] in {"success", "empty"} and coverage["status"] == "covered"
            if timestamps or valid_empty:
                folder = Path(group_names[panel.get("group") or "Ungrouped"])
                (staging / folder).mkdir(exist_ok=True)
                base = panel_filenames[panel["id"]]
                sheet_name = _sheet_name(panel["title"])
                column_chunks = [(i, min(i + max_columns - 1, len(series))) for i in range(0, len(series), max_columns - 1)] or [(0, 0)]
                row_chunks = [(i, min(i + max_rows - 1, len(timestamps))) for i in range(0, len(timestamps), max_rows - 1)] or [(0, 0)]
                for c_index, (c_start, c_end) in enumerate(column_chunks, 1):
                    for r_index, (r_start, r_end) in enumerate(row_chunks, 1):
                        suffix = f".r{r_index:03d}.c{c_index:03d}" if len(column_chunks) > 1 or len(row_chunks) > 1 else ""
                        relative = folder / f"{base}{suffix}.xlsx"
                        selected_times = timestamps[r_start:r_end]
                        selected_series = series[c_start:c_end]
                        _write_workbook(staging / relative, selected_series, legends[c_start:c_end], selected_times, sheet_name,
                                        [unit_plan[item["series_id"]] for item in selected_series])
                        entry["files"].append({
                            "path": (Path("exports") / generation / relative).as_posix(),
                            "sheet": sheet_name, "row_count": len(selected_times),
                            "columns": legends[c_start:c_end],
                            "series_ids": [s["series_id"] for s in series[c_start:c_end]],
                            "from_ms": selected_times[0] if selected_times else None,
                            "to_ms": selected_times[-1] if selected_times else None,
                        })
                entry["export_status"] = "data" if timestamps else "empty_selected_range"
            else:
                entry["export_status"] = "no_valid_display_data"
            manifest["panels"].append(entry)
        manifest["summary"] = dict(Counter(p["status"] for p in manifest["panels"]))
        # Publish files before the manifest. No existing generation is modified.
        os.replace(staging, destination)
        with temporary_manifest.open("w", encoding="utf-8") as output:
            json.dump(manifest, output, ensure_ascii=False, indent=2, allow_nan=False)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary_manifest, final_manifest)
        return final_manifest
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        # An unpublished generation is safe to remove; the old manifest remains.
        shutil.rmtree(destination, ignore_errors=True)
        temporary_manifest.unlink(missing_ok=True)
        raise
